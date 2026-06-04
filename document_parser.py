"""
多格式文档解析器。
"""
import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from html.parser import HTMLParser
from io import BytesIO


# ---------这是第2次更新，更新内容为：新增 TXT/PDF/DOC/DOCX/HTML/XLSX 文件解析能力---------
@dataclass
class ParsedDocument(object):
    text: str
    file_type: str
    metadata: dict


class SimpleHTMLTextParser(HTMLParser):
    """用标准库提取 HTML 可见文本，避免额外引入依赖。"""

    block_tags = {
        "address", "article", "aside", "blockquote", "br", "div", "dl",
        "fieldset", "figcaption", "figure", "footer", "form", "h1", "h2",
        "h3", "h4", "h5", "h6", "header", "hr", "li", "main", "nav", "ol",
        "p", "pre", "section", "table", "tbody", "td", "tfoot", "th",
        "thead", "tr", "ul",
    }
    ignore_tags = {"script", "style", "noscript"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.ignore_depth = 0

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag in self.ignore_tags:
            self.ignore_depth += 1
            return
        if tag in self.block_tags:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in self.ignore_tags:
            self.ignore_depth = max(0, self.ignore_depth - 1)
            return
        if tag in self.block_tags:
            self.parts.append("\n")

    def handle_data(self, data):
        if self.ignore_depth:
            return
        if data.strip():
            self.parts.append(data)

    def get_text(self):
        return "\n".join(self.parts)


def parse_uploaded_file(uploaded_file):
    """将 Streamlit UploadedFile 解析为可入库的纯文本。"""
    filename = uploaded_file.name
    file_type = _get_file_type(filename)
    file_bytes = uploaded_file.getvalue()

    parser_map = {
        "txt": parse_txt,
        "pdf": parse_pdf,
        "doc": parse_doc,
        "docx": parse_docx,
        "html": parse_html,
        "htm": parse_html,
        "xlsx": parse_xlsx,
    }
    parser = parser_map.get(file_type)
    if parser is None:
        raise ValueError(f"暂不支持的文件类型：{file_type}")

    parsed = parser(file_bytes, filename)
    if not parsed.text.strip():
        raise ValueError(f"{filename} 未解析到有效文本")
    return parsed


def parse_txt(file_bytes, filename):
    text, encoding = _decode_text(file_bytes)
    return ParsedDocument(
        text=_normalize_text(text),
        file_type="txt",
        metadata={
            "parser": "txt",
            "encoding": encoding,
        },
    )


def parse_pdf(file_bytes, filename):
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ImportError("解析 PDF 需要安装 pypdf，请先执行 pip install -r requirements.txt") from exc

    reader = PdfReader(BytesIO(file_bytes))
    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ValueError("PDF 已加密，当前无法解析") from exc

    content_lines = []
    for page_index, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        page_text = _normalize_text(page_text)
        if page_text:
            content_lines.append(f"[PDF 第 {page_index} 页]")
            content_lines.append(page_text)

    if not content_lines:
        parsed_text = ""
    else:
        parsed_text = _normalize_text("\n\n".join([f"文件：{filename}", "文件类型：PDF"] + content_lines))

    return ParsedDocument(
        text=parsed_text,
        file_type="pdf",
        metadata={
            "parser": "pypdf",
            "page_count": len(reader.pages),
        },
    )


def parse_docx(file_bytes, filename):
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise ImportError("解析 DOCX 需要安装 python-docx，请先执行 pip install -r requirements.txt") from exc

    doc = DocxDocument(BytesIO(file_bytes))
    content_lines = []

    for paragraph in doc.paragraphs:
        text = _normalize_inline_text(paragraph.text)
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name.lower().startswith("heading"):
            content_lines.append(f"# {text}")
        else:
            content_lines.append(text)

    for table_index, table in enumerate(doc.tables, start=1):
        rows = [
            [_normalize_inline_text(cell.text) for cell in row.cells]
            for row in table.rows
        ]
        content_lines.extend(_table_rows_to_text(rows, f"Word 表格 {table_index}"))

    if not content_lines:
        parsed_text = ""
    else:
        parsed_text = _normalize_text("\n\n".join([f"文件：{filename}", "文件类型：DOCX"] + content_lines))

    return ParsedDocument(
        text=parsed_text,
        file_type="docx",
        metadata={
            "parser": "python-docx",
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
        },
    )


def parse_doc(file_bytes, filename):
    """解析老版 .doc。优先使用 macOS textutil，其次使用 antiword。"""
    suffix = os.path.splitext(filename)[1] or ".doc"
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp_file:
            temp_file.write(file_bytes)
            temp_path = temp_file.name

        if shutil.which("textutil"):
            result = subprocess.run(
                ["textutil", "-convert", "txt", "-stdout", temp_path],
                check=True,
                capture_output=True,
            )
            text, encoding = _decode_text(result.stdout)
            parser_name = "textutil"
        elif shutil.which("antiword"):
            result = subprocess.run(
                ["antiword", temp_path],
                check=True,
                capture_output=True,
            )
            text, encoding = _decode_text(result.stdout)
            parser_name = "antiword"
        else:
            raise RuntimeError("解析 DOC 需要本机安装 textutil(macOS 自带) 或 antiword，推荐将 .doc 转为 .docx 后上传")
    finally:
        if temp_path and os.path.exists(temp_path):
            os.remove(temp_path)

    content_text = _normalize_text(text)
    parsed_text = _normalize_text(f"文件：{filename}\n文件类型：DOC\n\n{content_text}") if content_text else ""

    return ParsedDocument(
        text=parsed_text,
        file_type="doc",
        metadata={
            "parser": parser_name,
            "encoding": encoding,
        },
    )


def parse_html(file_bytes, filename):
    text, encoding = _decode_text(file_bytes)
    parser = SimpleHTMLTextParser()
    parser.feed(text)
    parser.close()
    content_text = _normalize_text(parser.get_text())
    parsed_text = _normalize_text(f"文件：{filename}\n文件类型：HTML\n\n{content_text}") if content_text else ""

    return ParsedDocument(
        text=parsed_text,
        file_type="html",
        metadata={
            "parser": "html.parser",
            "encoding": encoding,
        },
    )


def parse_xlsx(file_bytes, filename):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("解析 XLSX 需要安装 openpyxl，请先执行 pip install -r requirements.txt") from exc

    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    content_lines = []
    sheet_count = 0
    row_count = 0

    for worksheet in workbook.worksheets:
        sheet_count += 1
        sheet_lines = [f"[Excel 工作表：{worksheet.title}]"]
        header = None

        for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            values = [_cell_to_text(value) for value in row]
            if not any(values):
                continue

            if header is None:
                header = values
                sheet_lines.append(f"表头：{_join_values(values)}")
                continue

            row_count += 1
            row_text = _row_to_text(header, values)
            sheet_lines.append(f"工作表 {worksheet.title} 第 {row_index} 行：{row_text}")

        if header is not None:
            content_lines.extend(sheet_lines)

    workbook.close()
    if not content_lines:
        parsed_text = ""
    else:
        parsed_text = _normalize_text("\n\n".join([f"文件：{filename}", "文件类型：XLSX"] + content_lines))

    return ParsedDocument(
        text=parsed_text,
        file_type="xlsx",
        metadata={
            "parser": "openpyxl",
            "sheet_count": sheet_count,
            "data_row_count": row_count,
        },
    )


def _get_file_type(filename):
    return os.path.splitext(filename)[1].lower().lstrip(".")


def _decode_text(file_bytes):
    encodings = ("utf-8-sig", "utf-8", "gb18030", "gbk", "big5")
    for encoding in encodings:
        try:
            return file_bytes.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore"), "utf-8-ignore"


def _normalize_inline_text(text):
    return re.sub(r"\s+", " ", text or "").strip()


def _normalize_text(text):
    lines = [_normalize_inline_text(line) for line in (text or "").splitlines()]
    cleaned_lines = []
    last_blank = False
    for line in lines:
        if not line:
            if not last_blank:
                cleaned_lines.append("")
            last_blank = True
            continue
        cleaned_lines.append(line)
        last_blank = False
    return "\n".join(cleaned_lines).strip()


def _cell_to_text(value):
    if value is None:
        return ""
    return _normalize_inline_text(str(value))


def _join_values(values):
    return "；".join(value for value in values if value)


def _table_rows_to_text(rows, table_name):
    clean_rows = [row for row in rows if any(row)]
    if not clean_rows:
        return []

    lines = [f"[{table_name}]"]
    header = clean_rows[0]
    lines.append(f"{table_name} 表头：{_join_values(header)}")
    for row_index, row in enumerate(clean_rows[1:], start=1):
        lines.append(f"{table_name} 第 {row_index} 行：{_row_to_text(header, row)}")
    return lines


def _row_to_text(header, row):
    pairs = []
    max_length = max(len(header), len(row))
    for index in range(max_length):
        key = header[index] if index < len(header) and header[index] else f"列{index + 1}"
        value = row[index] if index < len(row) else ""
        if value:
            pairs.append(f"{key}：{value}")
    return "；".join(pairs)
# ---------第2次更新结束---------
